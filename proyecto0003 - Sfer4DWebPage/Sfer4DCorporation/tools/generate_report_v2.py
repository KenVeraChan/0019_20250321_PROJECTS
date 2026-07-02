from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from pathlib import Path
import textwrap

report_path = 'Sfer4D_report.docx'


def add_code_paragraph(doc, code):
    p = doc.add_paragraph()
    r = p.add_run(code)
    r.font.name = 'Courier New'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
    r.font.size = Pt(9)


def add_file_section(doc, file_path: Path, title: str, explanation: str, head_lines=0, tail_lines=0):
    doc.add_heading(title, level=3)
    doc.add_paragraph(explanation)
    try:
        text = file_path.read_text(encoding='utf-8')
    except Exception as e:
        add_code_paragraph(doc, f'Error leyendo {file_path}: {e}')
        return

    # Optionally trim very long files for readability
    lines = text.splitlines()
    if head_lines > 0 or tail_lines > 0:
        if tail_lines > 0:
            shown = lines[:head_lines] + (['...'] if len(lines) > head_lines + tail_lines else []) + lines[-tail_lines:]
        else:
            shown = lines[:head_lines]
        content = '\n'.join(shown)
    else:
        content = text

    add_code_paragraph(doc, content)


doc = Document()

doc.add_heading('Informe técnico: errores y correcciones (Sfer4D)', level=1)
doc.add_paragraph('Fecha: 2026-07-02')

doc.add_heading('Resumen ejecutivo', level=2)
doc.add_paragraph(
    'Este documento recoge los errores y problemas detectados durante la integración local y en Docker, '
    'junto con los cambios de código aplicados para resolverlos. Incluye fragmentos completos de los archivos modificados y explicaciones técnicas.'
)

doc.add_heading('1) Archivos modificados y fragmentos', level=2)

project_root = Path(__file__).resolve().parents[1]

# Add docker-compose.yml
add_file_section(
    doc,
    project_root / 'docker-compose.yml',
    'docker-compose.yml (completo)',
    'Se muestra el archivo completo que monta el directorio SQL y configura MySQL para ejecutar los scripts de inicialización.'
)

# backend/.env
add_file_section(
    doc,
    project_root / 'backend' / '.env',
    'backend/.env',
    'Variables de conexión que el backend usa para conectar al servidor MySQL. No es necesario incluir DB_NAME si se usan múltiples esquemas.'
)

# backend/db.js
add_file_section(
    doc,
    project_root / 'backend' / 'db.js',
    'backend/db.js (completo)',
    'Este archivo contiene la función de conexión con retry y el loader opcional que ejecuta los archivos .sql encontrados en `./sql` dentro del contenedor.'
)

# backend/server.js (fragment of interest)
add_file_section(
    doc,
    project_root / 'backend' / 'server.js',
    'backend/server.js (fragmento: areaConsultada y consulta)',
    'Se muestra la estructura `areaConsultada` y la consulta SQL que usa esquema.tabla para acceder a las tablas correctas entre las tres bases.',
    head_lines=1, tail_lines=200
)

# SQL dumps (show header and where separators were)
for sqlname in ['bbdd001_jefes_rrhh.sql', 'bbdd002_empleados.sql', 'bbdd003_clientes.sql']:
    add_file_section(
        doc,
        project_root / 'backend' / 'sql' / sqlname,
        f'backend/sql/{sqlname} (cabecera y primeras 40 líneas)',
        'Se muestra la cabecera del dump SQL y las primeras líneas donde originalmente había una línea con solo guiones que causaba error; ahora dichas líneas están comentadas con `--`.',
        head_lines=40
    )


doc.add_heading('2) Explicaciones y pasos aplicados', level=2)
doc.add_paragraph(textwrap.dedent('''
 - Se corrigió el montaje del directorio SQL en `docker-compose.yml` para apuntar a `./backend/sql`.
 - Se añadió un loader en `backend/db.js` que ejecuta los scripts SQL si están presentes.
 - Se comentaron las líneas de separador erróneas en los SQL dumps (líneas con `----`) usando `--` para convertirlas en comentarios SQL válidos.
 - Se documentó cómo realizar un reinicio completo del stack Docker y borrar el volumen (`docker compose down -v`) para forzar la ejecución de los scripts de inicialización.
'''))


doc.add_heading('3) Comandos útiles para verificación', level=2)
add_code_paragraph(doc, 'docker compose down -v\ndocker compose up --build\ndocker compose exec mysql mysql -uroot -p1234 -e "SHOW DATABASES;"')


doc.save(report_path)
print(f'Report written to {report_path}')
